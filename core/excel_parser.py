"""Unified Excel parser for LOI and BAIL data extraction."""

import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl

from core.formula_engine import resolve_formula
from core.inpi_client import INPIClient
from core.models import DossierData, InpiData, SocieteInfo
from config.settings import get_inpi_credentials, validate_inpi_credentials

logger = logging.getLogger(__name__)


class ExcelParser:
    """Parse Excel decision files and extract variables."""

    def __init__(self, source_path: str, config_path: str):
        """
        Args:
            source_path: Path to uploaded Excel file (Fiche de decision).
            config_path: Path to config Excel (Redaction LOI.xlsx or Redaction BAIL.xlsx).
        """
        self.source_path = Path(source_path)
        self.config_path = Path(config_path)
        self._effective_config_path = self._resolve_config_path()

    def _resolve_config_path(self) -> Path:
        """If the source file contains a 'Rédaction LOI' sheet, use it as config."""
        try:
            wb = openpyxl.load_workbook(str(self.source_path), read_only=True)
            sheet_names_lower = [s.lower() for s in wb.sheetnames]
            wb.close()
            if "rédaction loi" in sheet_names_lower or "redaction loi" in sheet_names_lower:
                logger.info(
                    "Source file contains 'Rédaction LOI' sheet — using it as config."
                )
                return self.source_path
        except Exception as e:
            logger.debug(f"Could not inspect source file for config sheet: {e}")
        return self.config_path

    def _format_cell_value(self, value) -> Optional[str]:
        """Format a cell value to string."""
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.strftime("%d/%m/%Y")
        if isinstance(value, (int, float)):
            return str(value)
        if isinstance(value, str):
            result = value.strip()
            return result if result else None
        return str(value)

    def _add_system_variables(self, variables: dict):
        """Add system-generated variables."""
        variables["Date d'aujourd'hui"] = datetime.now().strftime("%d/%m/%Y")

    def extract_variables(self) -> dict[str, str]:
        """Extract all variables from the source Excel using config mapping."""
        # Load source workbook (data_only=True to get cached formula values)
        source_wb = openpyxl.load_workbook(str(self.source_path), data_only=True)

        # Load config workbook in dual mode (using effective config path)
        config_wb_values = openpyxl.load_workbook(str(self._effective_config_path), data_only=True)
        config_wb_formulas = openpyxl.load_workbook(str(self._effective_config_path), data_only=False)

        # Find the config sheet: prefer "Rédaction LOI" if present, else first sheet
        config_sheet_name = None
        for name in config_wb_values.sheetnames:
            if name.lower() in ("rédaction loi", "redaction loi"):
                config_sheet_name = name
                break
        if not config_sheet_name:
            config_sheet_name = config_wb_values.sheetnames[0]

        ws_values = config_wb_values[config_sheet_name]
        ws_formulas = config_wb_formulas[config_sheet_name]

        variables: dict[str, str] = {}

        for row in range(2, ws_values.max_row + 1):
            nom = self._format_cell_value(ws_values.cell(row=row, column=1).value)
            if not nom:
                continue

            # Try data_only value first
            source_val = ws_values.cell(row=row, column=2).value
            # Fallback to formula if value is None or #REF!
            if source_val is None or (isinstance(source_val, str) and "#REF" in source_val):
                source_val = ws_formulas.cell(row=row, column=2).value

            source = self._format_cell_value(source_val) if source_val is not None else None

            if not source:
                continue

            if isinstance(source, str) and source.startswith("=") and "!" in source:
                # Cell reference formula -> resolve against source workbook
                value = resolve_formula(source, source_wb)
                if value:
                    variables[nom] = value
            elif isinstance(source, str) and ("[" in source and "]" in source):
                # Formula for later processing
                variables[f"_formula_{nom}"] = source
            else:
                # Direct value or description
                variables[nom] = source

        self._add_system_variables(variables)

        source_wb.close()
        config_wb_values.close()
        config_wb_formulas.close()

        return variables

    def extract_societe_info(self) -> dict[str, SocieteInfo]:
        """Extract Societe Bailleur info from config workbook.

        Looks for data in two places:
        1. A dedicated 'Société Bailleur' sheet (legacy format).
        2. Embedded in the 'Rédaction LOI' sheet (rows with a 'Société Bailleur'
           header row, typically row 35+, col A=name, B=header, D=footer).
        """
        config_wb = openpyxl.load_workbook(str(self._effective_config_path), data_only=True)

        result: dict[str, SocieteInfo] = {}

        # Strategy 1: dedicated sheet
        sheet_name = None
        for name in config_wb.sheetnames:
            if "bailleur" in name.lower() and ("societ" in name.lower() or "société" in name.lower()):
                sheet_name = name
                break

        if sheet_name:
            ws = config_wb[sheet_name]
            for row in range(2, ws.max_row + 1):
                nom = self._format_cell_value(ws.cell(row=row, column=1).value)
                if not nom:
                    continue
                header = self._format_cell_value(ws.cell(row=row, column=2).value) or nom
                footer = self._format_cell_value(ws.cell(row=row, column=3).value) or ""
                result[nom] = SocieteInfo(nom=nom, header_text=header, footer_text=footer)

        # Strategy 2: embedded in "Rédaction LOI" sheet (new unified format)
        if not result:
            redaction_sheet = None
            for name in config_wb.sheetnames:
                if name.lower() in ("rédaction loi", "redaction loi"):
                    redaction_sheet = name
                    break
            if redaction_sheet:
                ws = config_wb[redaction_sheet]
                # Find the "Société Bailleur" header row
                start_row = None
                for row in range(1, ws.max_row + 1):
                    val = self._format_cell_value(ws.cell(row=row, column=1).value)
                    if val and "bailleur" in val.lower() and ("societ" in val.lower() or "société" in val.lower()):
                        start_row = row + 1
                        break
                if start_row:
                    for row in range(start_row, ws.max_row + 1):
                        nom = self._format_cell_value(ws.cell(row=row, column=1).value)
                        if not nom:
                            break  # Empty row = end of société data
                        header = self._format_cell_value(ws.cell(row=row, column=2).value) or nom
                        # Footer is in column D (col 4) in the new format
                        footer = self._format_cell_value(ws.cell(row=row, column=4).value) or ""
                        result[nom] = SocieteInfo(nom=nom, header_text=header, footer_text=footer)

        config_wb.close()
        return result

    def enrich_from_inpi(self, variables: dict) -> Optional[InpiData]:
        """Enrich variables with INPI data if SIRET is available."""
        # Try multiple possible key names for SIRET
        siret = (
            variables.get("SIRET", "")
            or variables.get("N° DE SIRET", "")
            or variables.get("N° de SIRET", "")
            or variables.get("Numéro SIRET", "")
        )
        if not siret or not validate_inpi_credentials():
            return None

        try:
            creds = get_inpi_credentials()
            client = INPIClient(username=creds["username"], password=creds["password"])
            return client.get_company_info(siret)
        except Exception as e:
            logger.error(f"INPI enrichment error: {e}")
            return None

    def get_output_filename_loi(self, variables: dict) -> str:
        """Generate LOI output filename: 'YYYY MM DD - LOI NomPreneur.docx'."""
        nom_preneur = variables.get("Nom Preneur", "Preneur")
        date_loi = variables.get("Date LOI", "")
        try:
            dt = datetime.strptime(date_loi, "%d/%m/%Y")
            date_str = dt.strftime("%Y %m %d")
        except (ValueError, TypeError):
            date_str = datetime.now().strftime("%Y %m %d")
        return f"{date_str} - LOI {nom_preneur}.docx"

    def get_output_filename_bail(self, variables: dict) -> str:
        """Generate BAIL output filename: 'BAIL - NomPreneur - DateLOI.docx'."""
        nom_preneur = variables.get("Nom Preneur", "Preneur").replace("/", "-").replace("\\", "-")
        date_loi = variables.get("Date LOI", datetime.now().strftime("%d-%m-%Y"))
        # Replace / in date to avoid path issues
        date_loi = date_loi.replace("/", "-").replace("\\", "-")
        return f"BAIL - {nom_preneur} - {date_loi}.docx"
