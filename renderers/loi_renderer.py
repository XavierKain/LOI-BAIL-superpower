"""LOI Word document renderer: blue sections, headers/footers, special sections."""

import logging
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core.models import SocieteInfo
from generators.loi_generator import LOIGenerator
from renderers.word_engine import WordEngine

logger = logging.getLogger(__name__)


class LOIRenderer:
    """Render a LOI Word document from template + variables."""

    def __init__(self, template_path: str):
        self.template_path = Path(template_path)
        self.engine = WordEngine()

    def render(self, generator: LOIGenerator, societe_info: dict[str, SocieteInfo], output_path: str):
        """Generate the LOI document.

        Args:
            generator: LOIGenerator with all variables and clear list.
            societe_info: Dict of SocieteInfo keyed by societe name.
            output_path: Path to save the generated document.
        """
        doc = Document(str(self.template_path))
        variables = generator.get_all_variables()
        clear_list = generator.clear_list

        has_paliers = generator.has_palier_data()
        has_conditions = generator.has_conditions_suspensives()
        has_honoraires = generator.has_honoraires_preneurs()

        # Phase 1: Handle special sections
        to_delete = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()

            # Special: "Honoraires de commercialisation" (content-based deletion)
            if "Honoraires de commercialisation" in text and not has_honoraires:
                to_delete.append(paragraph)
                continue

            # Special: Blue "Remises" section
            if self.engine.is_paragraph_optional(paragraph) and "Remises" in text and "loyer" in text.lower():
                if not has_paliers:
                    to_delete.append(paragraph)
                    continue
                else:
                    # Keep and turn black
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(0, 0, 0)

            # Special: Blue "Condition(s) suspensive(s)" section
            if self.engine.is_paragraph_optional(paragraph) and "Condition" in text and "suspensive" in text.lower():
                if not has_conditions:
                    to_delete.append(paragraph)
                    continue
                else:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(0, 0, 0)
                    # Fall through to normal processing for [.] replacement

        # Phase 2: Standard placeholder replacement
        standard_deletes = self.engine.process_document_body(doc, variables, clear_list)
        to_delete.extend(standard_deletes)

        # Phase 3: Delete marked paragraphs
        self.engine.delete_paragraphs(to_delete)

        # Phase 4: Add borders to signature table cells
        self._add_signature_borders(doc)

        # Phase 5: Update headers/footers
        bailleur_name = (
            variables.get("Société Bailleur", "")
            or variables.get("Societe Bailleur", "")
        )
        if bailleur_name:
            # Try exact match first, then case-insensitive
            societe = societe_info.get(bailleur_name)
            if not societe:
                for key, info in societe_info.items():
                    if key.lower().strip() == bailleur_name.lower().strip():
                        societe = info
                        break
            if societe:
                self._update_headers_footers(doc, societe)

        # Save
        doc.save(output_path)
        logger.info(f"LOI document saved to {output_path}")

    def _update_headers_footers(self, document, societe: SocieteInfo):
        """Update document headers and footers for the societe bailleur."""
        for section in document.sections:
            section.top_margin = Inches(0.5)
            section.header_distance = Inches(0.3)
            section.bottom_margin = Inches(0.5)
            section.footer_distance = Inches(0.3)

            # Header
            header = section.header
            for p in list(header.paragraphs):
                p._element.getparent().remove(p._element)
            hp = header.add_paragraph()
            hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hr = hp.add_run(societe.header_text)
            hr.font.bold = True
            hr.font.size = Pt(22)
            hr.font.name = "Times New Roman"
            hp.paragraph_format.space_after = Pt(12)

            # Footer
            footer = section.footer
            for p in list(footer.paragraphs):
                p._element.getparent().remove(p._element)
            lines = societe.footer_text.split("\n")
            for idx, line in enumerate(lines):
                fp = footer.add_paragraph()
                fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                fr = fp.add_run(line)
                fr.font.size = Pt(9)
                fr.font.name = "Times New Roman"
                if idx == 0:
                    fp.paragraph_format.space_before = Pt(12)

    @staticmethod
    def _set_cell_border(cell, **kwargs):
        """Set borders on a table cell.

        Usage: _set_cell_border(cell, top={"sz": 4, "val": "single", "color": "000000"}, ...)
        Supported keys: top, bottom, left, right.
        """
        tc = cell._tc
        tc_pr = tc.get_or_add_tcPr()
        tc_borders = tc_pr.find(qn("w:tcBorders"))
        if tc_borders is None:
            tc_borders = OxmlElement("w:tcBorders")
            tc_pr.append(tc_borders)
        for edge, attrs in kwargs.items():
            element = OxmlElement(f"w:{edge}")
            for attr_name, attr_val in attrs.items():
                element.set(qn(f"w:{attr_name}"), str(attr_val))
            # Remove existing border element for this edge
            existing = tc_borders.find(qn(f"w:{edge}"))
            if existing is not None:
                tc_borders.remove(existing)
            tc_borders.append(element)

    def _add_signature_borders(self, document):
        """Add borders and space to signature table cells."""
        border_style = {"sz": "4", "val": "single", "color": "000000", "space": "0"}
        for table in document.tables:
            is_signature_table = False
            for row in table.rows:
                for cell in row.cells:
                    text_lower = cell.text.lower()
                    if "bailleur" in text_lower or "candidat" in text_lower:
                        is_signature_table = True
                        break
                if is_signature_table:
                    break
            if is_signature_table:
                for row in table.rows:
                    # Set minimum row height for signature space
                    tr = row._tr
                    tr_pr = tr.get_or_add_trPr()
                    tr_height = OxmlElement("w:trHeight")
                    tr_height.set(qn("w:val"), "3600")  # ~6.3cm height for signatures
                    tr_height.set(qn("w:hRule"), "atLeast")
                    tr_pr.append(tr_height)

                    for cell in row.cells:
                        self._set_cell_border(
                            cell,
                            top=border_style,
                            bottom=border_style,
                            left=border_style,
                            right=border_style,
                        )
