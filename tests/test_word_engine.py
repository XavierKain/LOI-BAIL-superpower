# tests/test_word_engine.py
import pytest
from docx import Document
from docx.shared import RGBColor, Pt
from renderers.word_engine import WordEngine


@pytest.fixture
def engine():
    return WordEngine()


def _make_paragraph(doc, text, color=None, bold=None, font_name=None, font_size=None):
    """Helper: create a paragraph with a single run."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.font.bold = bold
    if font_name:
        run.font.name = font_name
    if font_size:
        run.font.size = font_size
    return p


def _make_fragmented_paragraph(doc, fragments):
    """Helper: create a paragraph with multiple runs (fragmented placeholder).
    fragments = [("text", {optional font attrs}), ...]
    """
    p = doc.add_paragraph()
    for text, attrs in fragments:
        run = p.add_run(text)
        for k, v in attrs.items():
            if k == "bold":
                run.font.bold = v
            elif k == "name":
                run.font.name = v
            elif k == "size":
                run.font.size = v
    return p


def test_simple_replacement(engine):
    doc = Document()
    p = _make_paragraph(doc, "Le preneur [Nom Preneur] signe.")
    variables = {"Nom Preneur": "Test SAS"}
    result = engine.replace_placeholders(p, variables)
    assert result is None  # no deletion
    assert "Test SAS" in p.text
    assert "[Nom Preneur]" not in p.text


def test_missing_placeholder_turns_red(engine):
    doc = Document()
    p = _make_paragraph(doc, "Montant: [Montant du loyer]")
    variables = {}
    engine.replace_placeholders(p, variables)
    # The placeholder should still be there
    assert "[Montant du loyer]" in p.text
    # At least one run should be red
    has_red = any(
        run.font.color.rgb == RGBColor(255, 0, 0)
        for run in p.runs
        if run.font.color.rgb is not None
    )
    assert has_red


def test_fragmented_placeholder_replacement(engine):
    doc = Document()
    p = _make_fragmented_paragraph(doc, [
        ("Le preneur [Nom", {}),
        (" Preneur] signe.", {}),
    ])
    variables = {"Nom Preneur": "Test SAS"}
    engine.replace_placeholders(p, variables)
    assert "Test SAS" in p.text
    assert "[" not in p.text


def test_optional_blue_paragraph_deleted_when_no_data(engine):
    doc = Document()
    p = _make_paragraph(doc, "Optionnel: [Variable rare]", color=RGBColor(0, 0, 255))
    variables = {}
    result = engine.replace_placeholders(p, variables)
    assert result == "delete"


def test_optional_blue_paragraph_kept_when_data(engine):
    doc = Document()
    p = _make_paragraph(doc, "Optionnel: [Variable rare]", color=RGBColor(0, 0, 255))
    variables = {"Variable rare": "valeur"}
    result = engine.replace_placeholders(p, variables)
    assert result is None
    assert "valeur" in p.text
    # Runs should now be black (blue removed)
    for run in p.runs:
        if run.font.color.rgb is not None:
            assert run.font.color.rgb == RGBColor(0, 0, 0)


def test_clear_list_placeholder_replaced_with_empty(engine):
    doc = Document()
    p = _make_paragraph(doc, "President: [PRESIDENT DE LA SOCIETE]")
    variables = {"PRESIDENT DE LA SOCIETE": "Jean Dupont"}
    clear_list = ["PRESIDENT DE LA SOCIETE"]
    engine.replace_placeholders(p, variables, clear_list=clear_list)
    assert "Jean Dupont" not in p.text
    assert "[PRESIDENT DE LA SOCIETE]" not in p.text


def test_is_paragraph_optional_blue(engine):
    doc = Document()
    p = _make_paragraph(doc, "Blue text", color=RGBColor(0, 0, 200))
    assert engine.is_paragraph_optional(p) is True


def test_is_paragraph_optional_black(engine):
    doc = Document()
    p = _make_paragraph(doc, "Black text", color=RGBColor(0, 0, 0))
    assert engine.is_paragraph_optional(p) is False


def test_is_paragraph_optional_no_color(engine):
    doc = Document()
    p = _make_paragraph(doc, "No color")
    assert engine.is_paragraph_optional(p) is False


def test_format_preservation_simple(engine):
    doc = Document()
    p = _make_paragraph(doc, "Bold: [Nom]", bold=True, font_name="Arial", font_size=Pt(14))
    variables = {"Nom": "Test"}
    engine.replace_placeholders(p, variables)
    assert p.runs[0].font.bold is True
    assert p.runs[0].font.name == "Arial"


def test_en_lettres_placeholder(engine):
    doc = Document()
    p = _make_paragraph(doc, "Montant: [Montant du loyer en lettres]")
    variables = {"Montant du loyer": "160000"}
    engine.replace_placeholders(p, variables)
    assert "CENT SOIXANTE MILLE" in p.text


def test_multiple_placeholders_in_one_paragraph(engine):
    doc = Document()
    p = _make_paragraph(doc, "[Nom Preneur] signe le [Date LOI].")
    variables = {"Nom Preneur": "Test SAS", "Date LOI": "01/04/2026"}
    engine.replace_placeholders(p, variables)
    assert "Test SAS" in p.text
    assert "01/04/2026" in p.text


def test_process_document_body(engine):
    doc = Document()
    doc.add_paragraph("Hello [Nom]")
    doc.add_paragraph("World [Date]")
    variables = {"Nom": "Test", "Date": "2026"}
    engine.process_document_body(doc, variables)
    assert "Test" in doc.paragraphs[0].text
    assert "2026" in doc.paragraphs[1].text
