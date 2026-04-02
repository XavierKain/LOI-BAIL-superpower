"""End-to-end integration tests using real templates and data files."""
import pytest
import tempfile
from pathlib import Path
from docx import Document

from config.settings import TEMPLATE_LOI, TEMPLATE_BAIL, CONFIG_LOI, CONFIG_BAIL


@pytest.fixture
def loi_template_available():
    if not Path(TEMPLATE_LOI).exists():
        pytest.skip("Template LOI not found")


@pytest.fixture
def bail_template_available():
    if not Path(TEMPLATE_BAIL).exists():
        pytest.skip("Template BAIL not found")


def test_loi_template_has_placeholders(loi_template_available):
    """Verify the LOI template contains expected placeholders."""
    doc = Document(str(TEMPLATE_LOI))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "[" in full_text


def test_bail_template_has_article_placeholders(bail_template_available):
    """Verify the BAIL template contains {{ARTICLE}} placeholders."""
    doc = Document(str(TEMPLATE_BAIL))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "{{" in full_text


def test_word_engine_on_real_loi_template(loi_template_available):
    """Test WordEngine on the real LOI template with sample data."""
    from renderers.word_engine import WordEngine

    doc = Document(str(TEMPLATE_LOI))
    engine = WordEngine()

    variables = {
        "Nom Preneur": "TEST SAS",
        "Montant du loyer": "160 000",
        "Duree Bail": "9",
    }

    to_delete = engine.process_document_body(doc, variables)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "TEST SAS" in full_text


def test_full_loi_generation(loi_template_available):
    """Test complete LOI generation pipeline with mock data."""
    from core.models import DossierData
    from generators.shared import normaliser_noms_variables, calculer_variables_derivees
    from generators.loi_generator import LOIGenerator
    from renderers.loi_renderer import LOIRenderer

    variables = normaliser_noms_variables({
        "Nom Preneur": "INTEGRATION TEST SAS",
        "Montant du loyer": "160000",
        "Duree Bail": "9",
        "Numero et rue": "10 rue de la Paix",
        "Ville ou arrondissement": "Paris 2ème",
        "Date d'aujourd'hui": "01/04/2026",
        "Type Preneur": "SAS",
        "Societe Bailleur": "",
    })
    derivees = calculer_variables_derivees(variables, None)

    dossier = DossierData(
        variables=variables,
        variables_derivees=derivees,
        inpi_data=None,
        source_file=Path("/tmp/fake.xlsx"),
    )

    generator = LOIGenerator(dossier)
    renderer = LOIRenderer(str(TEMPLATE_LOI))

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        output_path = f.name

    renderer.render(generator, {}, output_path)

    assert Path(output_path).exists()
    doc = Document(output_path)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "INTEGRATION TEST SAS" in full_text
    assert "3/6/9" in full_text

    Path(output_path).unlink()


def test_bail_renderer_on_real_template(bail_template_available):
    """Test BAIL renderer on the real template with sample articles."""
    from core.models import ArticleResult
    from renderers.bail_renderer import BailRenderer

    articles = [
        ArticleResult(
            designation="Article 1",
            contenu="Le bail est conclu pour une durée de 9 ans.",
            placeholders_manquants=[],
        ),
    ]
    variables = {
        "Ville ou arrondissement": "Paris 2ème",
        "Date de signature": "22/04/2026",
    }

    renderer = BailRenderer(str(TEMPLATE_BAIL))

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        output_path = f.name

    renderer.render(articles, variables, output_path)

    assert Path(output_path).exists()
    doc = Document(output_path)
    assert len(doc.paragraphs) > 0

    Path(output_path).unlink()


def test_formatting_tags_parser():
    """Test HTML-like tag parsing used by BAIL renderer."""
    from renderers.bail_renderer import parse_formatting_tags

    segments = parse_formatting_tags("Normal <b>bold</b> end")
    assert len(segments) == 3
    assert segments[0] == ("Normal ", {})
    assert segments[1] == ("bold", {"bold": True})
    assert segments[2] == (" end", {})


def test_nested_formatting_tags():
    """Test nested HTML tags."""
    from renderers.bail_renderer import parse_formatting_tags

    segments = parse_formatting_tags("<b><i>bold italic</i></b>")
    assert len(segments) == 1
    assert segments[0] == ("bold italic", {"bold": True, "italic": True})
